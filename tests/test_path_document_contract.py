"""Native PDF paths and adaptive motion agree on document representation and static pixels."""

import io

import pytest
from PIL import Image, ImageChops
from pptx import Presentation
from quickthumb import AnimationSpec, Canvas, ExportPolicy, KeyframeSpec, ScaleTrack, TimingSpec
from quickthumb.errors import RenderingError

from tests._optional import require_cairosvg, require_pypdfium2
from tests.test_export_html import timelines
from tests.test_parent_documents import embedded_png


def test_transformed_path_document_contract(tmp_path):
    """Static transforms and motion use truthful document fallbacks with authored pixels."""
    pdfium, cairo = require_pypdfium2(), require_cairosvg()
    strict = ExportPolicy(unsupported_motion="error")
    data = "M2 20C2 2 38 2 38 20C38 38 2 38 2 20ZM16 15H24V25H16Z"
    # Given: both a static transform and a timed source require authored-static fallback
    animation = AnimationSpec.timeline(
        ScaleTrack(keyframes=[KeyframeSpec(time=0, value=2), KeyframeSpec(time=1, value=3)]),
        timing=TimingSpec(duration=1),
        easing="linear",
    )
    for animated in (False, True):
        source = Canvas(120, 100).path(
            data,
            (30, 30),
            40,
            40,
            "#E0406080",
            rotation=25,
            opacity=0.6,
            anchor=(0.25, 0.75),
            animation=animation if animated else None,
            id="fallback.path",
        )
        authored = source.sample().frames[0].to_image()
        bounds = authored.getbbox()
        assert bounds is not None
        expected = authored.crop(bounds)
        if animated:
            assert authored.tobytes() != source.render_frame(0).tobytes()
        white = Image.new("RGBA", authored.size, "white")
        white.alpha_composite(authored)

        # When: actual document files are exported, strict retry cannot replace them
        for format in ("svg", "html", "pdf", "pptx"):
            path = tmp_path / f"fallback-{animated}.{format}"
            result = source.export(path)
            payload = path.read_bytes()
            with pytest.raises(RenderingError) as error:
                source.export(path, policy=strict)
            assert path.read_bytes() == payload
            assert error.value.details[0].layer_id == "fallback.path"
            rows = [row for row in result.capability_report if row.feature == "path"]
            assert [(row.layer_id, row.support, row.fallback) for row in rows] == [
                ("fallback.path", "fallback", "rasterize")
            ]
            assert "static" in rows[0].message

            # Then: PDF's rendered page and embedded PNGs retain the canonical authored still
            if format == "pdf":
                with pdfium.PdfDocument(payload) as document:
                    page = document[0]
                    assert [obj.type for obj in page.get_objects()] == [3]
                    actual = page.render(scale=1).to_pil().convert("RGB")
                    difference = ImageChops.difference(actual, white.convert("RGB"))
                    assert max(difference.tobytes()) <= 2
            else:
                text = payload.decode() if format != "pptx" else payload
                image = embedded_png(text, format)
                assert image.size == expected.size and image.tobytes() == expected.tobytes()
                if format == "svg":
                    actual = Image.open(io.BytesIO(cairo.svg2png(bytestring=payload))).convert(
                        "RGBA"
                    )
                    assert actual.getbbox() == authored.getbbox()
                elif format == "html":
                    assert isinstance(text, str)
                    assert f"left:{bounds[0]}px;top:{bounds[1]}px;" in text
                    assert timelines(text) == [[]]
                else:
                    shape = Presentation(io.BytesIO(payload)).slides[0].shapes[0]
                    assert (shape.left, shape.top) == (bounds[0] * 9525, bounds[1] * 9525)
