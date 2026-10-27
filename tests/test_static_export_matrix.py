"""Static format checks share authored content, canonical samples and route expectations."""

import base64
import re
import xml.etree.ElementTree as ET
from io import BytesIO

import pytest
from PIL import Image
from quickthumb import Canvas, Deck, ExportPolicy
from quickthumb._svg_fidelity import _pixel_measurements, _visible_rgba, compare_pixels
from quickthumb.errors import RenderingError

from tests._optional import require_cairosvg, require_pypdfium2


@pytest.mark.parametrize("kind", ["canvas", "deck"])
@pytest.mark.parametrize("format", ["png", "svg", "html", "pdf", "pptx"])
def test_static_export_matrix(tmp_path, kind, format):
    """Static exports preserve shared content with explicit rendered or structural evidence."""
    # Given: an opaque mixed slide and a transparent slide, both without motion
    scenes = matrix_scenes(tmp_path)
    document = scenes[0] if kind == "canvas" else Deck(slides=scenes)
    original = document.to_json()
    (tmp_path / "source.json").write_text(original, encoding="utf-8")
    selected = scenes[:1] if kind == "canvas" else scenes
    for index, scene in enumerate(selected):
        scene.sample().frames[0].to_image().save(tmp_path / f"canonical-{index}.png")

    # When: the public exporter writes the requested format and optional observation
    if format == "pptx":
        pytest.importorskip("pptx")
    path = tmp_path / f"matrix.{format}"
    if format == "svg" and kind == "deck":
        with pytest.raises(RenderingError, match="SVG") as error:
            document.export(path)
        assert error.value.code == "unsupported_format" and not path.exists()
        results = [
            scene.export(tmp_path / f"slide-{index}.svg", policy=measured_policy("svg"))
            for index, scene in enumerate(selected)
        ]
    else:
        policy = measured_policy(format) if format in {"svg", "pdf"} and kind == "canvas" else None
        results = [document.export(path, policy=policy)]
    for index, result in enumerate(results):
        (tmp_path / f"receipt-{index}.json").write_text(result.model_dump_json(indent=2))

    # Then: each format is checked at its actual observation boundary, without a fake viewer verdict
    if format == "png":
        for scene, written in zip(selected, results[0].written_paths, strict=True):
            actual = Image.open(written).convert("RGBA")
            expected = scene.sample().frames[0].to_image()
            assert actual.size == expected.size and actual.tobytes() == expected.tobytes()
    elif format == "svg":
        for scene, result in zip(selected, results, strict=True):
            check_measurement(result, scene, "cairosvg")
            check_routes(result, [scene], deck=False)
            check_svg(result.written_paths[0], scene)
    elif format == "pdf":
        check_pdf(path, selected)
        if kind == "canvas":
            check_measurement(results[0], scenes[0], "pdfium")
            check_routes(results[0], selected, deck=False)
        else:
            assert results[0].pixel_metrics.fidelity is None
            assert not any(
                row.feature == "layer_representation" for row in results[0].capability_report
            )
    elif format == "pptx":
        check_pptx(path, selected)
        check_routes(results[0], selected, deck=kind == "deck")
        assert results[0].pixel_metrics.fidelity is None
    else:
        check_html(path, selected)
        check_routes(results[0], selected, deck=kind == "deck")
        assert results[0].pixel_metrics.fidelity is None
    assert document.to_json() == original


def test_deck_pdf_rejects_searchable_but_invisible_text(tmp_path):
    """A real Deck PDF with white searchable text fails the canonical text-region check."""
    require_pypdfium2()
    # Given: the reference has black text while the emitted native text is white on white
    scenes = matrix_scenes(tmp_path)
    wrong = Canvas(
        160,
        120,
        layers=[
            layer.model_copy(update={"color": "#FFFFFF"}) if layer.id == "text" else layer
            for layer in scenes[0].layers
        ],
    )

    # When: the real exporter retains searchable text and all other slide content
    path = tmp_path / "invisible-text.pdf"
    Deck(slides=[wrong, scenes[1]]).export(path)

    # Then: matching text objects cannot substitute for visible canonical glyphs
    with pytest.raises(AssertionError, match="text region"):
        check_pdf(path, scenes)


def matrix_scenes(tmp_path):
    image_path = tmp_path / "source.png"
    Image.new("RGB", (20, 16), "#2479D8").save(image_path)
    mixed = (
        Canvas(160, 120)
        .background(color="#FFFFFF", id="background")
        .text("Matrix", position=(8, 6), size=22, font="assets/fonts/Roboto-Regular.ttf", id="text")
        .shape("rectangle", (8, 45), 24, 20, "#EF3340", id="vector")
        .image(str(image_path), (50, 45), width=20, height=16, id="image")
        .shape(
            "rectangle",
            (100, 45),
            24,
            20,
            "#00A88F",
            id="mask",
            mask={"shape": "ellipse", "position": (100, 45), "width": 24, "height": 20},
        )
    )
    transparent = Canvas(160, 120).shape("rectangle", (8, 45), 24, 20, "#EF334080", id="alpha")
    return [mixed, transparent]


def measured_policy(format):
    if format == "svg":
        require_cairosvg()
    else:
        require_pypdfium2()
    return ExportPolicy(
        fidelity={
            format: {
                "pixel_tolerance": 0,
                "max_different_pixel_ratio": 0.08,
                "min_hash_similarity": 0.94,
            }
        }
    )


def check_measurement(result, scene, renderer):
    observation = result.pixel_metrics.fidelity
    assert observation is not None
    assert observation.renderer == renderer and observation.renderer_version
    assert (observation.actual_width, observation.actual_height) == (scene.width, scene.height)
    assert observation.canonical_environment == scene.sample().environment
    assert observation.verdict in {"exact", "tolerated"}
    assert observation.pixel_diff_ratio is not None and observation.pixel_diff_ratio <= 0.08
    assert observation.hash_similarity >= 0.94


def check_routes(result, scenes, *, deck):
    routes = [row for row in result.capability_report if row.feature == "layer_representation"]
    expected = []
    for index, scene in enumerate(scenes):
        for layer in scene.layers:
            support = "fallback" if layer.id in {"image", "mask"} else "native"
            expected.append((index if deck else None, layer.id, support))
    assert [(row.slide_index, row.layer_id, row.support) for row in routes] == expected
    assert all(row.native_attempt is (row.support == "native") for row in routes)
    assert {(row.slide_index, row.layer_id) for row in result.fallback_diagnostics} == {
        (index if deck else None, layer.id)
        for index, scene in enumerate(scenes)
        for layer in scene.layers
        if layer.id in {"image", "mask"}
    }
    assert all(row.reason for row in result.fallback_diagnostics)


def check_svg(path, scene):
    observed = Image.open(BytesIO(require_cairosvg().svg2png(url=path))).convert("RGBA")
    check_regions(observed, scene)
    root = ET.parse(path).getroot()
    namespace = {"s": "http://www.w3.org/2000/svg"}
    if len(scene.layers) == 1:
        assert len(root.findall(".//s:rect", namespace)) == 1
        return
    assert "Matrix" in "".join(root.itertext())
    assert len(root.findall(".//s:rect", namespace)) == 2
    images = root.findall(".//s:image", namespace)
    assert len(images) == 2
    payload = images[1].attrib["{http://www.w3.org/1999/xlink}href"].split(",", 1)[1]
    actual = Image.open(BytesIO(base64.b64decode(payload))).convert("RGBA")
    expected = mask_pixels(scene)
    assert actual.size == expected.size and actual.tobytes() == expected.tobytes()


def check_pdf(path, scenes):
    pdfium = require_pypdfium2()
    with pdfium.PdfDocument(path) as document:
        assert len(document) == len(scenes)
        for index, scene in enumerate(scenes):
            page = document[index]
            try:
                assert page.get_size() == (scene.width, scene.height)
                types = [obj.type for obj in page.get_objects()]
                assert pdfium.raw.FPDF_PAGEOBJ_PATH in types
                if index == 0:
                    assert pdfium.raw.FPDF_PAGEOBJ_TEXT in types
                    assert types.count(pdfium.raw.FPDF_PAGEOBJ_IMAGE) == 2
                    text = page.get_textpage()
                    try:
                        assert "Matrix" in text.get_text_range()
                    finally:
                        text.close()
                bitmap = page.render(scale=1, fill_color=(0, 0, 0, 0))
                try:
                    actual = bitmap.to_pil().convert("RGBA")
                    check_regions(actual, scene)
                finally:
                    bitmap.close()
            finally:
                page.close()


def check_pptx(path, scenes):
    pptx = pytest.importorskip("pptx")
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    document = pptx.Presentation(path)
    assert len(document.slides) == len(scenes)
    assert (document.slide_width, document.slide_height) == (160 * 9525, 120 * 9525)
    for index, slide in enumerate(document.slides):
        text_boxes = [shape for shape in slide.shapes if shape.has_text_frame and shape.text]
        pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if index == 0:
            assert [shape.text for shape in text_boxes] == ["Matrix"]
            assert len(pictures) == 2
            mask = pictures[1]
            expected = BytesIO()
            mask_pixels(scenes[0]).save(expected, format="PNG")
            assert mask.image.blob == expected.getvalue()
            assert (mask.left, mask.top, mask.width, mask.height) == tuple(
                value * 9525 for value in (100, 45, 24, 20)
            )
        else:
            assert not text_boxes and not pictures
        native_shapes = [
            shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        ]
        assert len(native_shapes) == (2 if index == 0 else 1)


def check_html(path, scenes):
    markup = path.read_text(encoding="utf-8")
    assert markup.count('class="qt-stage"') == len(scenes)
    assert "Matrix" in markup and "background:rgb(239,51,64)" in markup
    images = re.findall(r'<img[^>]*src="data:image/png;base64,([^"]+)"', markup)
    assert len(images) == 2
    actual = Image.open(BytesIO(base64.b64decode(images[1]))).convert("RGBA")
    expected = mask_pixels(scenes[0])
    assert actual.size == expected.size and actual.tobytes() == expected.tobytes()


def mask_pixels(scene):
    mask = Canvas(scene.width, scene.height, layers=[scene.layers[-1]])
    return mask.sample().frames[0].to_image().crop((100, 45, 124, 65))


def check_regions(actual, scene):
    actual = _visible_rgba(actual)
    expected = _visible_rgba(scene.sample().frames[0].to_image())
    assert actual.size == expected.size
    assert actual.getpixel((150, 110)) == expected.getpixel((150, 110))
    if len(scene.layers) > 1:
        box = (8, 6, 78, 34)
        text = compare_pixels(
            expected.crop(box),
            actual.crop(box),
            {
                "pixel_tolerance": 16,
                "max_different_pixel_ratio": 0.30,
                "min_hash_similarity": 0.95,
            },
        )
        assert text["verdict"] in {"exact", "tolerated"}, f"text region: {text}"
    # Local color/placement checks prevent small missing components passing a whole-page budget.
    for box in [(8, 45, 32, 65), (50, 45, 70, 61), (108, 51, 116, 59)]:
        observed, reference = actual.crop(box), expected.crop(box)
        if len(scene.layers) == 1:
            assert observed.getchannel("A").tobytes() == reference.getchannel("A").tobytes()
            assert _pixel_measurements(reference, observed, 2)["pixel_diff_ratio"] == 0
        else:
            assert observed.tobytes() == reference.tobytes()
